"""Kardex persistence, valuation, barcode and Inventory integration regression."""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _product(
    client: TestClient,
    *,
    sku: str = "SKU-001",
    barcode: str | None = "2000000000015",
    method: str = "WEIGHTED_AVERAGE",
    initial_stock: float | None = None,
    initial_cost: float | None = None,
    warehouse: str = "Lima",
) -> dict:
    payload = {
        "sku": sku,
        "barcode": barcode,
        "barcode_type": "EAN13" if barcode else None,
        "name": f"Producto {sku}",
        "category": "Core",
        "unit_of_measure": "unidad",
        "warehouse": warehouse,
        "valuation_method": method,
        "minimum_stock": 20,
    }
    if initial_stock is not None:
        payload["initial_stock"] = initial_stock
        payload["initial_unit_cost"] = initial_cost
    response = client.post("/api/v1/kardex/products", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _movement(
    client: TestClient,
    product_id: str,
    movement_type: str,
    quantity: float,
    *,
    cost: float | None = None,
    day: int = 1,
) -> dict:
    response = client.post(
        "/api/v1/kardex/movements",
        json={
            "product_id": product_id,
            "timestamp": (datetime.now(UTC) + timedelta(seconds=day)).isoformat(),
            "movement_type": movement_type,
            "document_reference": f"DOC-{day}-{movement_type}",
            "quantity": quantity,
            "unit_cost": cost,
            "source": "manual",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _forecast(client: TestClient, *, product: str = "NX-I01", location: str = "Lima") -> dict:
    start = date(2024, 1, 1)
    rows = ["date,product,category,location,demand"]
    for index in range(100):
        rows.append(f"{start + timedelta(days=index)},{product},Core,{location},{20 + index % 7}")
    dataset = client.post(
        "/api/v1/datasets/upload",
        files={"file": ("kardex.csv", ("\n".join(rows) + "\n").encode(), "text/csv")},
    ).json()
    client.post(f"/api/v1/datasets/{dataset['id']}/validate")
    client.post(f"/api/v1/datasets/{dataset['id']}/ready")
    response = client.post(
        "/api/v1/forecast-runs",
        json={
            "dataset_id": dataset["id"],
            "product": product,
            "category": "Core",
            "location": location,
            "frequency": "daily",
            "horizon": 14,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_product_creation_initial_balance_and_lookup(client: TestClient) -> None:
    product = _product(client, initial_stock=100, initial_cost=10)
    assert product["sku"] == "SKU-001"
    assert product["stock_quantity"] == 100
    assert client.get(
        "/api/v1/kardex/products/lookup", params={"identifier": product["barcode"]}
    ).json()["id"] == product["id"]
    assert client.get(
        "/api/v1/kardex/products/lookup", params={"identifier": "sku-001"}
    ).json()["id"] == product["id"]


@pytest.mark.parametrize("field", ["sku", "barcode"])
def test_sku_and_barcode_are_unique(client: TestClient, field: str) -> None:
    _product(client)
    response = client.post(
        "/api/v1/kardex/products",
        json={
            "sku": "SKU-001" if field == "sku" else "SKU-002",
            "barcode": "2000000000015" if field == "barcode" else "2000000000022",
            "barcode_type": "EAN13",
            "name": "Duplicado",
        },
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "kardex_product_conflict"


def test_invalid_barcode_checksum_is_rejected(client: TestClient) -> None:
    response = client.post(
        "/api/v1/kardex/products",
        json={
            "sku": "BAD-001",
            "barcode": "2000000000014",
            "barcode_type": "EAN13",
            "name": "Inválido",
        },
    )
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "kardex_invalid_barcode"


def test_weighted_average_matches_documented_example(client: TestClient) -> None:
    product = _product(client, initial_stock=100, initial_cost=10)
    inbound = _movement(client, product["id"], "PURCHASE_IN", 50, cost=12, day=2)
    assert inbound["balance_quantity"] == 150
    assert inbound["balance_total"] == 1600
    assert inbound["balance_unit_cost"] == pytest.approx(10.666667)
    outbound = _movement(client, product["id"], "SALE_OUT", 30, day=3)
    assert outbound["unit_cost"] == pytest.approx(10.666667)
    assert outbound["total_cost"] == 320
    assert outbound["balance_quantity"] == 120
    assert outbound["balance_unit_cost"] == pytest.approx(10.666667)


def test_weighted_average_multiple_entries_and_adjustments(client: TestClient) -> None:
    product = _product(client, initial_stock=10, initial_cost=5)
    _movement(client, product["id"], "PURCHASE_IN", 10, cost=7, day=2)
    plus = _movement(client, product["id"], "POSITIVE_ADJUSTMENT", 5, cost=6, day=3)
    minus = _movement(client, product["id"], "NEGATIVE_ADJUSTMENT", 3, day=4)
    assert plus["balance_quantity"] == 25
    assert plus["balance_unit_cost"] == 6
    assert minus["balance_quantity"] == 22
    assert minus["balance_total"] == 132


def test_fifo_consumes_multiple_layers_and_persists_allocations(client: TestClient) -> None:
    product = _product(client, method="FIFO", initial_stock=100, initial_cost=10)
    _movement(client, product["id"], "PURCHASE_IN", 50, cost=12, day=2)
    outbound = _movement(client, product["id"], "SALE_OUT", 120, day=3)
    assert outbound["total_cost"] == 1240
    assert outbound["balance_quantity"] == 30
    assert outbound["balance_unit_cost"] == 12
    assert outbound["balance_total"] == 360
    assert [(item["quantity"], item["unit_cost"]) for item in outbound["allocations"]] == [
        (100, 10),
        (20, 12),
    ]
    recovered = client.get(f"/api/v1/kardex/movements/{outbound['id']}").json()
    assert recovered["allocations"] == outbound["allocations"]


def test_fifo_partial_layer_remains_reconstructible(client: TestClient) -> None:
    product = _product(client, method="FIFO", initial_stock=40, initial_cost=3)
    outbound = _movement(client, product["id"], "SALE_OUT", 15, day=2)
    assert outbound["total_cost"] == 45
    assert outbound["balance_quantity"] == 25
    assert outbound["balance_total"] == 75


def test_negative_stock_and_backdated_movements_are_rejected(client: TestClient) -> None:
    product = _product(client, initial_stock=10, initial_cost=4)
    response = client.post(
        "/api/v1/kardex/movements",
        json={
            "product_id": product["id"],
            "timestamp": (datetime.now(UTC) + timedelta(seconds=2)).isoformat(),
            "movement_type": "SALE_OUT",
            "quantity": 11,
        },
    )
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "kardex_negative_stock"
    backdated = client.post(
        "/api/v1/kardex/movements",
        json={
            "product_id": product["id"],
            "timestamp": datetime(2025, 12, 31, tzinfo=UTC).isoformat(),
            "movement_type": "PURCHASE_IN",
            "quantity": 1,
            "unit_cost": 4,
        },
    )
    assert backdated.status_code == 400
    assert backdated.json()["error"]["code"] == "kardex_backdated_movement"


def test_preview_does_not_persist(client: TestClient) -> None:
    product = _product(client, initial_stock=10, initial_cost=4)
    payload = {
        "product_id": product["id"],
        "timestamp": (datetime.now(UTC) + timedelta(seconds=2)).isoformat(),
        "movement_type": "SALE_OUT",
        "quantity": 2,
    }
    preview = client.post("/api/v1/kardex/movements/preview", json=payload)
    assert preview.status_code == 200
    assert preview.json()["resulting_quantity"] == 8
    assert len(client.get(f"/api/v1/kardex/products/{product['id']}/movements").json()) == 1


def test_kardex_is_returned_in_reverse_chronological_order(client: TestClient) -> None:
    product = _product(client, initial_stock=10, initial_cost=4)
    _movement(client, product["id"], "PURCHASE_IN", 2, cost=5, day=2)
    rows = client.get(f"/api/v1/kardex/products/{product['id']}/movements").json()
    assert rows[0]["timestamp"] > rows[1]["timestamp"]


def test_no_movements_is_missing_while_confirmed_zero_is_real(client: TestClient) -> None:
    empty = _product(client, sku="EMPTY", barcode=None)
    assert client.get(f"/api/v1/kardex/products/{empty['id']}/balance").json()["quantity"] is None
    zero = _product(client, sku="ZERO", barcode="2000000000022", initial_stock=5, initial_cost=1)
    _movement(client, zero["id"], "SALE_OUT", 5, day=2)
    balance = client.get(f"/api/v1/kardex/products/{zero['id']}/balance").json()
    assert balance["has_kardex"] is True
    assert balance["quantity"] == 0


def test_summary_uses_derived_balances(client: TestClient) -> None:
    _product(client, initial_stock=10, initial_cost=4)
    summary = client.get("/api/v1/kardex/summary").json()
    assert summary["active_skus"] == 1
    assert summary["units_in_stock"] == 10
    assert summary["inventory_value"] == 40
    assert summary["below_minimum"] == 1


def test_demo_is_deterministic_and_safe_to_repeat(client: TestClient) -> None:
    first = client.post("/api/v1/kardex/demo/regenerate")
    second = client.post("/api/v1/kardex/demo/regenerate")
    assert first.status_code == second.status_code == 200
    assert [item["id"] for item in first.json()["products"]] == [
        item["id"] for item in second.json()["products"]
    ]
    beverage = next(item for item in second.json()["products"] if item["sku"] == "BEB-001")
    assert beverage["stock_quantity"] == 120
    assert beverage["stock_value"] == pytest.approx(1280)


def test_export_is_a_real_excel_workbook(client: TestClient) -> None:
    product = _product(client, initial_stock=10, initial_cost=4)
    response = client.get("/api/v1/kardex/export.xlsx", params={"product_id": product["id"]})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Kardex", "Resumen"]
    assert workbook["Kardex"]["A1"].value == "Fecha"
    assert workbook["Kardex"].max_row == 2


def test_inventory_can_explicitly_use_kardex_balance(client: TestClient) -> None:
    forecast = _forecast(client)
    product = _product(
        client,
        sku="NX-I01",
        barcode="2000000000022",
        initial_stock=77,
        initial_cost=9,
    )
    cutoff = (datetime.now(UTC) + timedelta(minutes=5)).isoformat()
    response = client.post(
        "/api/v1/inventory/preflight",
        json={
            "forecast_run_id": forecast["id"],
            "cutoff": cutoff,
            "inventory_source": "kardex",
            "kardex_product_id": product["id"],
            "operational_inputs": {},
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["inventory_source"] == "kardex"
    assert response.json()["inventory_source_detail"]["quantity"] == 77
    run = client.post(
        "/api/v1/inventory",
        json={
            "forecast_run_id": forecast["id"],
            "cutoff": cutoff,
            "inventory_source": "kardex",
            "kardex_product_id": product["id"],
            "operational_inputs": {},
        },
    )
    assert run.status_code == 201, run.text
    assert run.json()["items"][0]["inventory_on_hand"] == 77
    assert run.json()["items"][0]["inputs"]["inventory_on_hand"]["source_type"] == "kardex"


def test_inventory_kardex_absence_is_not_zero_and_manual_mode_is_unchanged(
    client: TestClient,
) -> None:
    forecast = _forecast(client)
    product = _product(client, sku="NX-I01", barcode="2000000000022")
    cutoff = (datetime.now(UTC) + timedelta(minutes=5)).isoformat()
    kardex = client.post(
        "/api/v1/inventory/preflight",
        json={
            "forecast_run_id": forecast["id"],
            "cutoff": cutoff,
            "inventory_source": "kardex",
            "kardex_product_id": product["id"],
            "operational_inputs": {},
        },
    ).json()
    assert "inventory_on_hand" in kardex["missing_inputs"]
    assert kardex["inventory_source_detail"]["has_kardex"] is False
    manual = client.post(
        "/api/v1/inventory/preflight",
        json={
            "forecast_run_id": forecast["id"],
            "cutoff": cutoff,
            "operational_inputs": {"inventory_on_hand": {"value": 11, "status": "available"}},
        },
    ).json()
    assert manual["inventory_source"] == "manual"
    assert manual["inventory_source_detail"] is None
    assert "inventory_on_hand" not in manual["missing_inputs"]


def test_kardex_scope_mismatch_and_source_conflict_are_explicit(client: TestClient) -> None:
    forecast = _forecast(client)
    wrong = _product(client, sku="OTHER", barcode="2000000000022", initial_stock=5, initial_cost=1)
    cutoff = (datetime.now(UTC) + timedelta(minutes=5)).isoformat()
    mismatch = client.post(
        "/api/v1/inventory/preflight",
        json={
            "forecast_run_id": forecast["id"],
            "cutoff": cutoff,
            "inventory_source": "kardex",
            "kardex_product_id": wrong["id"],
        },
    )
    assert mismatch.status_code == 400
    assert mismatch.json()["error"]["code"] == "inventory_kardex_incompatible"
    conflict = client.post(
        "/api/v1/inventory/preflight",
        json={
            "forecast_run_id": forecast["id"],
            "cutoff": cutoff,
            "inventory_source": "kardex",
            "kardex_product_id": wrong["id"],
            "operational_inputs": {
                "inventory_on_hand": {"value": 99, "status": "available"}
            },
        },
    )
    assert conflict.status_code == 400
    assert conflict.json()["error"]["code"] == "inventory_stock_source_conflict"
