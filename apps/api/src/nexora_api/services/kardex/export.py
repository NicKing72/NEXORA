"""Backend-rendered XLSX export of the exact persisted ledger."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from nexora_api.models.kardex import InventoryProduct
from nexora_api.services.kardex.balances import current_values
from nexora_api.services.kardex.movements import list_movements

HEADERS = [
    "Fecha",
    "Documento",
    "SKU",
    "Código de barras",
    "Producto",
    "Movimiento",
    "Entrada Cantidad",
    "Entrada Costo Unitario",
    "Entrada Total",
    "Salida Cantidad",
    "Salida Costo Unitario",
    "Salida Total",
    "Saldo Cantidad",
    "Saldo Costo",
    "Saldo Total",
    "Método de Valorización",
    "Observación",
]


def export_kardex_xlsx(db: Session, product_id: str | None = None) -> bytes:
    products = db.query(InventoryProduct).order_by(InventoryProduct.sku).all()
    product_map = {product.id: product for product in products}
    movements = list(reversed(list_movements(db, product_id=product_id, limit=10000)))
    workbook = Workbook()
    ledger = workbook.active
    ledger.title = "Kardex"
    ledger.append(HEADERS)
    for cell in ledger[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="213735")
    for movement in movements:
        product = product_map[movement.product_id]
        is_in = movement.direction == "in"
        ledger.append(
            [
                movement.timestamp.replace(tzinfo=None),
                movement.document_reference,
                product.sku,
                product.barcode,
                product.name,
                movement.movement_type,
                float(movement.quantity) if is_in else None,
                float(movement.unit_cost) if is_in else None,
                float(movement.total_cost) if is_in else None,
                float(movement.quantity) if not is_in else None,
                float(movement.unit_cost) if not is_in else None,
                float(movement.total_cost) if not is_in else None,
                float(movement.balance_quantity),
                float(movement.balance_unit_cost),
                float(movement.balance_total),
                movement.valuation_method,
                movement.note,
            ]
        )
    ledger.freeze_panes = "A2"
    ledger.auto_filter.ref = ledger.dimensions
    for column in ledger.columns:
        ledger.column_dimensions[column[0].column_letter].width = min(
            38, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    summary = workbook.create_sheet("Resumen")
    summary.append(["SKU", "Producto", "Código", "Stock", "Costo", "Valor", "Método"])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="213735")
    selected = products if product_id is None else [product_map[product_id]]
    for product in selected:
        quantity, unit_cost, value, _ = current_values(db, product.id)
        summary.append(
            [
                product.sku,
                product.name,
                product.barcode,
                None if quantity is None else float(quantity),
                None if unit_cost is None else float(unit_cost),
                None if value is None else float(value),
                product.valuation_method,
            ]
        )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
